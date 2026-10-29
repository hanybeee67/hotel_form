# -*- coding: utf-8 -*-
"""변환 로직 테스트 — 모든 데이터는 가상 데이터(실제 투숙객·카드 정보 아님)."""
import datetime as dt
import json
import os
import sys

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, Side

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import converters as cv  # noqa: E402

BDATE = "2026-10-22"
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


@pytest.fixture
def cfg():
    with open(os.path.join(ROOT, "config.json"), encoding="utf-8") as f:
        return json.load(f)


@pytest.fixture(autouse=True)
def no_templates(monkeypatch, tmp_path):
    """기본은 양식 파일이 없는 상태로 테스트 (저장소의 templates/ 영향 받지 않게)."""
    monkeypatch.setattr(cv, "app_dir", lambda: str(tmp_path / "app"))
    monkeypatch.setattr(cv, "bundled_dir", lambda: str(tmp_path / "app"))
    monkeypatch.setattr(cv, "_settings_path", lambda: str(tmp_path / "settings" / "settings.json"))
    return tmp_path / "app"


def _xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.save(path)
    return str(path)


def _values(path):
    ws = load_workbook(path).active
    return ws, [[c.value for c in row] for row in ws.iter_rows()]


INCOME = [
    ["매출구분", "투숙객명", "도착일", "출발일", "객실타입", "거래처", "객실#", "회원번호", "RATE", "DC", "입력자",
     "입력일시", "객실수", "NET", "봉사료", "부가세", "예수금", "매출", "입금"],
    ["ROOM", "가상손님B", "2026-10-22", "2026-10-23", "D/T", "가상여행", "0305", "", 105000, 6, "", "", 1, 90000, 0,
     9000, 0, 99000, 0],
    ["ROOM", "가상손님A", "2026-10-22", "2026-10-23", "F/T", "가상회사", "0302", "", 125000, 0, "", "", 1, 80000, 0,
     8000, 0, 88000, 0],
    ["ROOM", "가상단체", "2026-10-22", "2026-10-23", "D/T", "가상회사", "1030", "", 0, 0, "", "", 1, 0, 0,
     0, 0, 50000, 0],
]
CARD = [
    ["영업일자", "순번", "빌번호[업장]", "Acct#[객실]", "구분", "카드번호", "할부기간", "부가세", "승인금액", "실정산금액",
     "사업자번호", "응답", "발급사", "매입사", "가맹점번호", "승인일시", "승인번호", "거래번호"],
    [BDATE, 1, 0, "A001", "카드승인", "1111-xxxx-xxxx-xxxx", 0, 9000, 99000, 99000, "0000000000", "성공", "가상발급",
     "가상카드", "123", "", "11111111", "1SN"],
    [BDATE, 2, 0, "", "카드승인", "", 0, 9000, 99000, 0, "0000000000", "실패", "", "", "", "", "", ""],
    [BDATE, 3, 0, "A002", "현금영수증", "", 0, 1000, 11000, 11000, "0000000000", "성공", "", "", "", "", "222", ""],
    [None, 3, None, None, None, None, None, None, 209000, None, None, None, None, None, None, None, None, None],
]
CREDIT = [
    ["지점명", "업장명", "빌번호", "객실번호", "입금종류", "입금일자", "고객명", "도착일", "출발일", "거래처", "입금액", "카드번호"],
    ["가상호텔", "ROOM", 1, "0609", "대외후불", BDATE, "가상손님C", BDATE, "2026-10-23", "가상여행", 150000, ""],
    ["가상호텔", "ROOM", 1, "0313", "대외후불", BDATE, "가상손님D", BDATE, "2026-10-23", "가상예약", 80402, ""],
    ["계", 2, None, None, None, None, None, None, None, None, 230402, None],
]
BREAKFAST = [
    ["순번", "예약자", "연락처", "비고", "고객명", "객실번호", "인원수[성인]", "인원수[소인]"],
    [1, "가", "", "2인 조식 포함", "가상손님E", "0503", 2, 0],
    [2, "나", "", "룸온리", "가상손님F", "0601", 1, 0],
    [3, "다", "", "조포", "가상손님G", "0412", 1, 1],
    [4, "라", "", "1인 조포", "가상악단", "0601", 1, 0],
    [5, "라", "", "1인 조포", "가상악단", "0602", 1, 0],
    [6, "라", "", "1인 조포", "가상악단", "0702", 1, 0],
    [7, "마", "", "1인 조포", "가상그룹-Kim", "0703", 1, 0],
    [8, "마", "", "2인 조포", "가상그룹-Lee", "0704", 2, 0],
    [9, "바", "", "조포", "가상단체계정", "1287", 1, 0],
]


def test_income(cfg, tmp_path):
    src = _xlsx(tmp_path / "수입.xlsx", INCOME)
    res = cv.convert_file("income", src, str(tmp_path / "out"), BDATE, cfg)
    ws, v = _values(res["output"])
    assert v[0][0] == "수 입 일 보 (10/22)"
    assert [r[0] for r in v[2:5]] == ["0302", "0305", "1030"]  # 객실 순 정렬, 앞자리 0 유지
    assert v[2][7] == 88000  # 매출 = 부가세 포함 금액 (NET 아님)
    assert v[2][8] == "=H3*G3"  # 누계 = 매출 × 숙박수 (호텔 실제 일보, 자기 자신을 가리키면 순환 참조)
    assert v[5][6] == "계" and v[5][7] == "=SUM(H3:H5)" and v[5][8] == "=SUM(I3:I5)"
    assert ws["F3"].fill.fgColor.rgb.endswith("FFF2B3")  # 선/후 노란 칸
    dvs = {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}
    assert dvs["F3:F5"].formula1 == '"카.선,현.선,후"'  # 선/후는 목록에서 고르는 노란 칸 (직접 써도 됨)
    assert not dvs["F3:F5"].showErrorMessage
    assert [v[2][6], v[3][6]] == [1, 1]  # 숙박수 = 출발일 - 도착일 (자동)
    assert not ws["G3"].fill.fgColor.rgb.endswith("FFF2B3")
    assert ws.page_setup.blackAndWhite  # 노란 칸은 인쇄되지 않게
    assert ws["B3"].font.name == "굴림" and ws["B3"].font.sz == 18 and ws["A1"].font.sz == 28
    assert os.path.basename(res["output"]) == "수입일보 2026-10-22.xlsx"
    # 수입일보는 매출이 잡힌 줄은 객실 범위 밖(1030 외부 조식 등)도 모두 집계 (2026-10-28)
    assert v[4][0] == "1030" and v[4][7] == 50000
    assert not any("범위" in w for w in res["warnings"])
    # 표 전체에 검은 선 (바깥 굵게, 안쪽 가늘게), A4 너비 맞춤
    for ref in ("A2", "E3", "I6", "B4"):
        b = ws[ref].border
        assert all(getattr(b, s).style for s in ("left", "right", "top", "bottom")), ref
        assert b.left.color.rgb == "FF000000"  # 불투명 검정
    assert ws["A3"].border.left.style == "medium" and ws["I3"].border.right.style == "medium"
    assert ws["E6"].border.bottom.style == "medium" and ws["E4"].border.bottom.style == "thin"
    # 노란 입력 칸(선/후·숙박수) 둘레는 굵은 선이라 서로 맞닿아도 경계가 보인다
    assert ws["F3"].border.right.style == "medium" and ws["G3"].border.left.style == "medium"
    assert ws["F3"].border.left.style == "medium" and ws["E3"].border.right.style == "medium"
    assert ws["F3"].fill.fgColor.rgb == "FFFFF2B3"
    ps = ws.page_setup
    assert int(ps.paperSize) == 9 and ps.orientation == "landscape" and ps.fitToWidth == 1 and ps.fitToHeight == 2  # A4 너비 1장 × 최대 2장
    assert ws.print_options.horizontalCentered and ws.print_title_rows == "$1:$2"


def test_card(cfg, tmp_path):
    src = _xlsx(tmp_path / "카드.xlsx", CARD)
    res = cv.convert_file("card", src, str(tmp_path / "out"), BDATE, cfg)
    _, v = _values(res["output"])
    assert len(v) == 9  # 제목, 머리글, 성공 1건, 한벽루 빈 줄 4, 현금영수증, 계
    assert v[2][1] == "1111-XXXX-XXXX-XXXX"
    assert v[2][3] == "가상카드"  # 발급사가 아니라 매입사
    assert v[7][5] == "현금영수증" and v[7][7] == 11000
    assert v[8][6] == "계" and v[8][7] == "=SUM(H3:H7)+H8"
    ws = load_workbook(res["output"]).active
    assert ws["A3"].fill.fgColor.rgb.endswith("FFF2B3")  # 객실은 Acct#를 보고 직접 입력하는 노란 칸


def test_credit(cfg, tmp_path):
    src = _xlsx(tmp_path / "후불.xlsx", CREDIT)
    res = cv.convert_file("credit", src, str(tmp_path / "out"), BDATE, cfg)
    _, v = _values(res["output"])
    assert v[0][0] == "대외후불리스트(10/22)"
    assert [r[0] for r in v[2:4]] == ["0313", "0609"]
    assert v[4][4] == "계" and v[4][5] == "=SUM(F3:F4)"
    assert not any("합계" in w and "다릅니다" in w for w in res["warnings"])  # 원본 합계와 일치
    assert os.path.basename(res["output"]) == "10월22일 대외후불리스트.xlsx"


def test_breakfast(cfg, tmp_path):
    src = _xlsx(tmp_path / "inhouse.xlsx", BREAKFAST)
    res = cv.convert_file("breakfast", src, str(tmp_path / "out"), BDATE, cfg)
    ws, v = _values(res["output"])
    rows = [r for r in v[2:] if r[0] and r[0].isdigit()]
    # 룸온리 제외, 인원 = N인 또는 성인+소인. 단체도 묶지 않고 모든 방을 객실 순서대로 한 줄씩 (2026-10-26 변경)
    assert [r[0] for r in rows] == ["0412", "0503", "0601", "0602", "0702", "0703", "0704"]
    assert [r[2] for r in rows] == [2, 2, 1, 1, 1, 1, 2]
    assert not any(r[0] == "단체" for r in v)
    assert not any(r[0] == "1287" for r in v)
    assert ws["B3"].font.name == "맑은 고딕" and ws["B3"].font.sz < 18  # A4 너비에 맞게 줄임


def test_wrong_file_message(cfg, tmp_path):
    src = _xlsx(tmp_path / "아무거나.xlsx", [["가", "나"], [1, 2]])
    with pytest.raises(cv.ConversionError) as e:
        cv.convert_file("credit", src, str(tmp_path / "out"), BDATE, cfg)
    assert "필수 항목" in str(e.value)


def test_detect_type(cfg, tmp_path):
    assert cv.detect_type(_xlsx(tmp_path / "a.xlsx", CARD), cfg) == "card"
    assert cv.detect_type(_xlsx(tmp_path / "b.xlsx", CREDIT), cfg) == "credit"
    assert cv.detect_type(_xlsx(tmp_path / "c.xlsx", BREAKFAST), cfg) == "breakfast"
    assert cv.detect_type(_xlsx(tmp_path / "d.xlsx", INCOME), cfg) == "income"


def test_template_keeps_hotel_format(cfg, tmp_path, no_templates):
    """가상 호텔 양식을 넣으면 양식의 폰트·행 높이·열 너비·병합·인쇄 방향이 그대로 유지된다."""
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "대외후불리스트(10/01)"
    ws["A1"].font = Font(name="굴림", size=16)
    ws.merge_cells("A1:F1")
    for i, h in enumerate(["객실", "고객명", "도착일", "출발일", "거래처", "금액"], 1):
        ws.cell(2, i, h).border = box
        ws.cell(3, i).font = Font(name="굴림", size=10)
        ws.cell(3, i).border = box
        ws.cell(4, i).border = box
    ws["E4"] = "계"
    ws.column_dimensions["B"].width = 9.5
    ws.row_dimensions[3].height = 21.75
    ws.page_setup.orientation = "landscape"
    os.makedirs(no_templates / "templates")
    wb.save(no_templates / "templates" / "대외후불리스트.xlsx")

    src = _xlsx(tmp_path / "후불.xlsx", CREDIT)
    res = cv.convert_file("credit", src, str(tmp_path / "out"), BDATE, cfg)
    assert not any("양식 파일" in w for w in res["warnings"])
    ws, v = _values(res["output"])
    assert v[0][0] == "대외후불리스트(10/22)" and ws["A1"].font.sz == 16
    assert ws["B3"].font.sz == 10 and ws.row_dimensions[4].height == 21.75
    assert ws.column_dimensions["B"].width == 9.5
    assert ws.page_setup.orientation == "landscape"
    assert v[4][4] == "계" and v[4][5] == "=SUM(F3:F4)"


@pytest.mark.parametrize("kind,data,title_size,body_size,font", [
    ("income", INCOME, 28, 18, "굴림"),
    ("card", CARD, 26, 16, "굴림"),
    ("credit", CREDIT, 28, 16, "굴림"),
    ("breakfast", BREAKFAST, 26, 18, "맑은 고딕"),
])
def test_repo_templates(cfg, tmp_path, monkeypatch, kind, data, title_size, body_size, font):
    """저장소의 호텔 빈 양식(templates/)으로 변환하면 호텔 양식의 글꼴·크기·열 너비가 그대로 나온다."""
    monkeypatch.setattr(cv, "app_dir", lambda: ROOT)
    tpl = load_workbook(os.path.join(ROOT, cfg["converters"][kind]["template"])).active
    src = _xlsx(tmp_path / f"{kind}.xlsx", data)
    res = cv.convert_file(kind, src, str(tmp_path / "out"), BDATE, cfg)
    assert not any("양식 파일" in w for w in res["warnings"])
    ws, v = _values(res["output"])
    # 조식명단은 A4 너비에 맞게 열 너비·글자 크기를 같은 비율로 조절 (2026-10-27), 나머지는 양식 그대로
    f = ws.column_dimensions["B"].width / tpl.column_dimensions["B"].width if kind == "breakfast" else 1
    near = lambda a, b: abs(a - b) <= 0.5
    assert ws["A1"].font.name == font and near(ws["A1"].font.sz, title_size * f)
    assert ws["B3"].font.name == font and near(ws["B3"].font.sz, body_size * f) and ws["B3"].font.b
    for col in "ABCD":
        assert abs(ws.column_dimensions[col].width - tpl.column_dimensions[col].width * f) < 0.01
    if kind == "breakfast":  # 조식명단은 A4를 꽉 채우도록 행 높이를 늘린다 (2026-10-27)
        assert ws.row_dimensions[3].height >= tpl.row_dimensions[3].height * f
    else:
        assert ws.row_dimensions[3].height == tpl.row_dimensions[3].height
    import datetime as dt
    expect = dt.date.today().strftime("%m/%d") if kind == "breakfast" else "10/22"  # 조식은 서류 만드는 날
    assert expect in v[0][0]
    # 깨끗한 표준 서식만: 옛 .xls 색상표·흰색 채우기 없음, 표의 모든 칸 네 변에 불투명 검정 선
    import re
    import zipfile
    styles = zipfile.ZipFile(res["output"]).read("xl/styles.xml").decode()
    # 호텔 PC 프로그램은 '테두리 적용' 표시가 없으면 선을 그리지 않는다 (빌드 16에서 확인됨) → 모든 셀 서식에 필수
    xfs = re.findall(r"<xf [^>]*>", re.search(r"<cellXfs.*?</cellXfs>", styles, re.S).group(0))
    assert xfs and all('applyBorder="1"' in x and 'applyFill="1"' in x and 'applyFont="1"' in x for x in xfs)
    fills = re.search(r"<fills.*?</fills>", styles, re.S).group(0)
    assert "FFFFFF\"" not in fills  # 흰색 채우기 없음 (노란 입력 칸만)
    last_row = max(r for r in range(1, ws.max_row + 1) if any(ws.cell(r, c).value is not None for c in range(1, 4)))
    right = {"income": 9, "card": 8, "credit": 6, "breakfast": 5}[kind]
    for r in range(2, last_row + 1):
        for c in range(1, right + 1):
            b = ws.cell(r, c).border
            merged = next((m for m in ws.merged_cells.ranges if (r, c) in m.cells), None)
            for s in ("left", "right", "top", "bottom"):
                if merged and {"left": c > merged.min_col, "right": c < merged.max_col,
                               "top": r > merged.min_row, "bottom": r < merged.max_row}[s]:
                    continue  # 병합 칸 안쪽 변은 보이지 않으므로 건너뜀
                sd = getattr(b, s)
                assert sd.style in ("thin", "medium") and sd.color.rgb == "FF000000", (r, c, s)
    assert ws["B3"].border.bottom.style and ws["B3"].border.bottom.color.rgb == "FF000000"
    assert ws["B3"].fill.fill_type in (None, "none")  # 흰색 채우기 없음
    if kind in ("income", "card", "credit"):  # 금액은 세 자리마다 콤마
        money = {"income": "H", "card": "H", "credit": "F"}[kind]
        nums = [c for c in ws[money][2:] if isinstance(c.value, (int, float)) or str(c.value).startswith("=")]
        assert nums and all("#,##0" in c.number_format for c in nums), [c.number_format for c in nums]
    if kind == "breakfast":
        firsts = [r[0] for r in v]
        assert "외부식사" not in firsts and any(str(x).startswith("조식 명단에") for x in firsts)
        total = next(r for r in v if r[0] == "합계")  # 외부식사 대신 합계 행 + 인원 합 수식
        n = len([r for r in v[2:] if r[0] and (str(r[0]).isdigit() or r[0] == "단체")])
        assert total[2] == f"=SUM(C3:C{2 + n})"
        assert ws["E3"].border.right.style == "medium"  # 병합된 Sign 칸의 오른쪽 열도 테두리 유지
        assert ws.auto_filter.ref.startswith("A2:E")


def test_output_folder_is_program_folder(cfg, tmp_path, monkeypatch):
    """결과는 원본 옆이 아니라 프로그램(exe) 폴더 안의 '마감결과'에 모인다 (한희 폴더)."""
    prog = tmp_path / "한희"
    monkeypatch.setattr(cv, "app_dir", lambda: str(prog))
    assert cv.default_out_dir(str(tmp_path / "다운로드" / "erp.xlsx"), cfg) == str(prog / "마감결과")
    assert cv.default_out_dir(None, {**cfg, "output_dir": "결과모음"}) == str(prog / "결과모음")
    assert cv.default_out_dir(None, {**cfg, "output_dir": str(tmp_path / "X")}) == str(tmp_path / "X")


def test_unwritable_program_folder_falls_back(cfg, tmp_path, monkeypatch):
    """공용폴더처럼 프로그램 폴더에 쓸 수 없으면 내 PC 폴더(문서\\마감결과 등)에 저장하고 알려 준다."""
    alt = tmp_path / "문서" / "마감결과"
    real = cv._can_write
    monkeypatch.setattr(cv, "_can_write", lambda folder: False if "공용" in str(folder) else real(folder))
    monkeypatch.setattr(cv, "fallback_dir", lambda: str(alt))
    src = _xlsx(tmp_path / "후불.xlsx", CREDIT)
    res = cv.convert_file("credit", src, str(tmp_path / "공용" / "한희" / "마감결과"), BDATE, cfg)
    assert os.path.dirname(res["output"]) == str(alt) and os.path.exists(res["output"])
    assert any("권한" in w for w in res["warnings"])


def test_error_log(tmp_path, monkeypatch):
    monkeypatch.setattr(cv, "app_dir", lambda: str(tmp_path))
    try:
        raise ValueError("시험 오류")
    except ValueError:
        path = cv.write_error_log("시험")
    with open(path, encoding="utf-8") as f:
        assert path and "ValueError" in f.read()


def test_chosen_save_folder_wins(cfg, tmp_path, monkeypatch):
    """바탕화면에서 실행해도 '저장 위치 변경'으로 고른 한희 폴더에 결과가 모인다 (설정은 내 PC에 기억)."""
    monkeypatch.setattr(cv, "app_dir", lambda: str(tmp_path / "바탕화면"))
    assert cv.default_out_dir(None, cfg) == str(tmp_path / "바탕화면" / "마감결과")
    cv.save_user_setting("output_dir", str(tmp_path / "공용" / "한희" / "마감결과"))
    assert cv.load_user_settings()["output_dir"].endswith("마감결과")
    assert cv.default_out_dir(None, cfg) == str(tmp_path / "공용" / "한희" / "마감결과")


@pytest.mark.parametrize("raw,masked", [
    ("홍길동", "홍0동"), ("김철", "김0"), ("남궁민수", "남00수"), ("정*영", "정0영"),
    ("진양그룹-Noah Raymond", "진양그룹-Noah Raymond"), ("CHRISTOPH", "CHRISTOPH"),
    ("Zhu bao jun / Lin jia hao", "Zhu bao jun / Lin jia hao"), ("코원필오케스트라", "코원필오케스트라"), ("", ""),
    # 4글자는 두 글자 성일 때만 사람 이름 — 회사 이름은 그대로 (2026-10-27 '씨티소재' → '씨00재' 문제)
    ("씨티소재", "씨티소재"), ("한빛여행", "한빛여행"), ("삼성전자", "삼성전자"), ("황보영희", "황00희"),
    ("제갈공명", "제00명"), ("김영희철", "김영희철"), ("대한항공", "대한항공"), ("하나투어", "하나투어"),
])
def test_mask_korean_name(raw, masked):
    assert cv.mask_korean_name(raw) == masked


def test_names_masked_in_documents(cfg, tmp_path):
    rows = [CREDIT[0], CREDIT[1][:6] + ["홍길동"] + CREDIT[1][7:], CREDIT[2][:6] + ["Kim Minsu"] + CREDIT[2][7:]]
    src = _xlsx(tmp_path / "후불.xlsx", rows)
    _, v = _values(cv.convert_file("credit", src, str(tmp_path / "out"), BDATE, cfg)["output"])
    assert sorted(r[1] for r in v[2:4]) == ["Kim Minsu", "홍길동"]  # 대외후불·수입일보는 이름 그대로 (2026-10-26)
    bf = [BREAKFAST[0], [1, "가", "", "1인 조포", "홍길동", "0503", 1, 0], [2, "나", "", "1인 조포", "Kim Minsu", "0504", 1, 0]]
    _, v = _values(cv.convert_file("breakfast", _xlsx(tmp_path / "bf.xlsx", bf), str(tmp_path / "out"), BDATE, cfg)["output"])
    assert [r[1] for r in v[2:4]] == ["홍0동", "Kim Minsu"]  # 조식명단에만 가운데 0


def test_many_rows_fit_two_a4_pages(cfg, tmp_path, monkeypatch):
    """객실이 많아도 A4 2장에 들어가게 행 높이를 줄이고, 글자 크기는 그대로."""
    monkeypatch.setattr(cv, "app_dir", lambda: ROOT)  # 실제 호텔 양식 사용
    rows = [INCOME[0]] + [["ROOM", f"손님{i}", BDATE, "2026-10-23", "D/T", "회사", f"{300 + (i % 13) + 1 + (i // 13) * 100:04d}",
                           "", 0, 0, "", "", 1, 0, 0, 0, 0, 99000, 0] for i in range(65)]
    src = _xlsx(tmp_path / "수입.xlsx", rows)
    ws, v = _values(cv.convert_file("income", src, str(tmp_path / "out"), BDATE, cfg)["output"])
    assert ws.row_dimensions[3].height < 39 and ws["B3"].font.sz == 18
    assert ws.page_setup.fitToHeight == 2
    # 높이 합이 A4 가로 2장(너비 맞춤 배율 기준) 안에 들어가는지 계산으로 확인
    total_w = sum(cv._col_points(ws.column_dimensions[c].width) for c in "ABCDEFGHI")
    scale = min(1, (842 - 0.8 * 72) / total_w)
    used = sum(ws.row_dimensions[r].height or 15 for r in range(3, ws.max_row + 1))
    budget = 2 * ((595 - 72) / scale - (57 + 39))
    # 행 높이를 줄여 들어가거나, 글자가 잘리지 않는 최소 높이(18pt×1.3)까지 줄였으면 나머지는 엑셀의 2장 맞춤이 축소
    assert used <= budget + 1 or all(ws.row_dimensions[r].height <= 18 * 1.3 + 0.01 for r in range(3, 60))


@pytest.mark.parametrize("remark,ok,pax", [
    ("1인 조포", True, 1), ("2인 조식 포함", True, 2), ("조식포함", True, None), ("조식 2인", True, 2),
    ("2인실 1인 조포", True, 1), ("조식추가 1명", True, 1), ("B/F 포함", True, None),
    ("조식 포함\n\n현장 결제\n\n2인 1실 조식 추가 비용 13,000원", True, None),  # 추가요금 안내의 '2인'은 인원 아님
    ("룸온니\n\n대외후불", False, None), ("가예약\n\n조식 포함\n\n현장 결제", True, None),
    ("룸온리", False, None), ("룸 온리", False, None), ("RO", False, None), ("Room Only", False, None),
    ("조식불포함", False, None), ("조식 미포함", False, None), ("조식X", False, None), ("조식 없음", False, None),
    ("노조식", False, None), ("", False, None), ("늦은 체크인", False, None),
])
def test_breakfast_remark_rules(cfg, remark, ok, pax):
    """실제 조식 47명인데 59명으로 집계된 문제(2026-10-26): '조식불포함' 같은 표현이 조식으로 잡히면 안 된다."""
    got_ok, got_pax, _ = cv._bf_people(remark, cfg["converters"]["breakfast"])
    assert (got_ok, got_pax) == (ok, pax)


def test_breakfast_same_room_counted_once(cfg, tmp_path):
    rows = [BREAKFAST[0],
            [1, "가", "", "2인 조식 포함", "홍길동", "0503", 2, 0],
            [2, "가", "", "2인 조식 포함", "홍길순", "0503", 2, 0],   # 같은 방 동숙자 줄
            [3, "나", "", "조식불포함", "김철수", "0504", 1, 0],
            [4, "다", "", "룸온리", "이영희", "0505", 1, 0]]
    src = _xlsx(tmp_path / "inhouse.xlsx", rows)
    res = cv.convert_file("breakfast", src, str(tmp_path / "out"), BDATE, cfg)
    _, v = _values(res["output"])
    rooms = [r[0] for r in v[2:] if r[0] and str(r[0]).isdigit()]
    assert rooms == ["0503"]
    assert any("1객실 2명" in w for w in res["warnings"]) and any("한 번만" in w for w in res["warnings"])


def test_income_sections_like_hotel_report(cfg, tmp_path):
    """수입일보 = 호텔 실제 일보 모양 (2026-10-28): 매출구분 묶음마다 계(앞 묶음까지 합산), 매출 0·소계 줄 제외,
    ROOM은 객실마다 한 줄 · 숙박수 = 이번 달 묵은 날 · 누계 = 매출×숙박수, 다른 묶음은 이름(구분) · 누계 = 매출."""
    def row(room, sales, kind="ROOM", arr="2026-10-26", dep="2026-10-30", guest="손님"):
        return [kind, guest, arr, dep, "D/T", "회사", room, "", 0, 0, "", "", 1, 0, 0, 0, 0, sales, 0]
    tot = lambda label, amt: [label] + [""] * 16 + [amt, 0]
    rows = [INCOME[0], row("0303", 100000), row("0302", 0), row("0302", 99000, arr="2026-08-25"),
            row("0304", 0), tot("ROOM 소계", 199000),
            row("0302", 5000, "Extra-Charge"), tot("Extra-Charge 소계", 5000),
            row("1030", 13000, "조식", arr="2019-06-30", dep="2027-12-31", guest="외부"), tot("조식 소계", 13000),
            row("0303", 0, "카드"), tot("카드 소계", 0), tot("계", 217000)]
    res = cv.convert_file("income", _xlsx(tmp_path / "수입.xlsx", rows), str(tmp_path / "out"), "2026-10-28", cfg)
    assert not any("원본 합계" in w for w in res["warnings"])  # 소계를 두 번 세지 않음 → ERP 계와 같음
    _, v = _values(res["output"])
    body = [r[:9] for r in v[2:]]
    assert [r[0] for r in body] == ["0302", "0303", None, "0302", None, "1030", None]
    assert body[0][6] == 28 and body[1][6] == 3  # 8/25 도착 장기 → 10/1부터 28, 10/26 도착 → 3
    assert body[0][8] == "=H3*G3"
    assert body[2][6:9] == ["계", "=SUM(H3:H4)", "=SUM(I3:I4)"]
    assert body[3][1] == "손님(EXTRA CHARGE)" and body[3][6] is None and body[3][8] == "=H6"
    assert body[4][6:9] == ["계", "=H5+SUM(H6:H6)", "=I5+SUM(I6:I6)"]
    assert body[5][1] == "외부(조식)" and body[6][7] == "=H7+SUM(H8:H8)"
    assert any("0인 줄" in w for w in res["warnings"])


@pytest.mark.parametrize("arr,dep,bd,n", [("2026-10-26", "2026-10-29", "2026-10-28", 3),
                                          ("2026-08-25", "2026-11-13", "2026-10-28", 28),
                                          ("2026-10-28", "2026-10-29", "2026-10-28", 1),
                                          ("", "2026-10-29", "2026-10-28", None)])
def test_nights_so_far(arr, dep, bd, n):
    assert cv._nights_so_far(arr, dep, dt.date.fromisoformat(bd)) == n


@pytest.mark.parametrize("arr,dep,n", [("2026-10-22", "2026-10-23", 1), ("2026-10-11", "2026-10-23", 12),
                                       ("2026-10-22", "2026-10-22", 0), ("", "2026-10-23", None), ("2026-10-23", "2026-10-22", None)])
def test_nights(arr, dep, n):
    assert cv._nights(arr, dep) == n


def _bf_fill_ratio(ws):
    """조식명단 본문이 N쪽 공간을 얼마나 채우는지 (1.0 = 꽉 참)"""
    m = ws.page_margins
    total_w = sum(cv._col_points(ws.column_dimensions[c].width) for c in "ABCDE")
    scale = min(1, (595 - (m.left + m.right) * 72) / total_w)
    repeat = sum(ws.row_dimensions[r].height for r in (1, 2))
    per_page = (842 - (m.top + m.bottom) * 72) / scale - repeat
    body = sum(ws.row_dimensions[r].height or 15 for r in range(3, ws.max_row + 1))
    return body / (ws.page_setup.fitToHeight * per_page)


@pytest.mark.parametrize("n_rooms,pages", [(3, 1), (12, 1), (25, 2), (45, 2), (80, 2)])
def test_breakfast_fills_a4(cfg, tmp_path, monkeypatch, n_rooms, pages):
    """조식명단만: 1장 분량이면 1장, 넘치면 2장을 꽉 채운다 (최대 2장, 2026-10-27)."""
    monkeypatch.setattr(cv, "app_dir", lambda: ROOT)
    rooms = [f"0{f}{n:02d}" for f in range(3, 8) for n in range(1, 14)]
    rows = [BREAKFAST[0]] + [[i, "가", "", "1인 조포", f"손님{i}", rooms[i % 65], 1, 0] for i in range(n_rooms)]
    res = cv.convert_file("breakfast", _xlsx(tmp_path / "bf.xlsx", rows), str(tmp_path / "out"), BDATE, cfg)
    ws = load_workbook(res["output"]).active
    assert ws.page_setup.fitToHeight == pages
    ratio = _bf_fill_ratio(ws)
    if n_rooms <= 45:
        assert 0.9 <= ratio <= 1.0, ratio  # 꽉 참 (넘치지 않음)
    # 가로도 A4 너비에 맞춤: 넉넉한 width_px 계산으로도 너비 안 (2026-10-27 Sign 열 잘림 수정)
    usable_w = 595 - (ws.page_margins.left + ws.page_margins.right) * 72
    total_w = sum(cv._col_points(ws.column_dimensions[c].width, cfg["converters"]["breakfast"]["width_px"]) for c in "ABCDE")
    assert 0.95 * usable_w <= total_w <= usable_w, (total_w, usable_w)
    assert ws["B3"].font.sz < 18 and ws["B2"].font.sz < 18  # 열과 같은 비율로 글자도 줄임
    assert ws.page_setup.scale in (None, 100)


def test_other_documents_not_stretched(cfg, tmp_path, monkeypatch):
    """조식명단 외 서류는 늘리지 않는다 (짧으면 양식 높이 그대로)."""
    monkeypatch.setattr(cv, "app_dir", lambda: ROOT)
    res = cv.convert_file("credit", _xlsx(tmp_path / "c.xlsx", CREDIT), str(tmp_path / "out"), BDATE, cfg)
    ws = load_workbook(res["output"]).active
    assert ws.row_dimensions[3].height == 39.95


@pytest.mark.parametrize("n_rooms", [20, 30, 43, 55])
def test_breakfast_two_pages_same_size(cfg, tmp_path, monkeypatch, n_rooms):
    """조식명단 2장: 두 장을 나란히 놓으므로 표 크기가 똑같고 맨 아래 줄이 딱 맞아야 한다 (2026-10-27)."""
    monkeypatch.setattr(cv, "app_dir", lambda: ROOT)
    rooms = [f"0{f}{n:02d}" for f in range(3, 8) for n in range(1, 14)]
    rows = [BREAKFAST[0]] + [[i, "가", "", "1인 조포", "손님", rooms[i], 1, 0] for i in range(n_rooms)]
    res = cv.convert_file("breakfast", _xlsx(tmp_path / "bf.xlsx", rows), str(tmp_path / "out"), BDATE, cfg)
    ws = load_workbook(res["output"]).active
    breaks = [b.id for b in ws.row_breaks.brk]
    assert len(breaks) == 1 and ws.sheet_properties.pageSetUpPr.fitToPage is False and ws.page_setup.scale
    cut, last = breaks[0], ws.max_row
    page1 = sum(ws.row_dimensions[r].height for r in range(3, cut + 1))
    page2 = sum(ws.row_dimensions[r].height for r in range(cut + 1, last + 1))
    assert abs(page1 - page2) < 0.5, (page1, page2)
    assert ws.cell(last, 1).value.startswith("조식 명단에")  # 안내 문구가 2쪽 맨 아래


# ---------------------------------------------------------------- 조식 인원수 (내일 아침, 2026-10-27)
INHOUSE = [
    ["순번", "비고", "고객명", "객실번호", "도착일", "출발일", "인원수[성인]", "인원수[소인]"],
    [1, "2인 조식 포함", "가상손님A", "0301", "2026-10-21", "2026-10-24", 2, 0],
    [2, "룸온리", "가상손님B", "0302", "2026-10-22", "2026-10-23", 1, 0],
    [3, "조포", "가상손님C", "0303", "2026-10-22", "2026-10-23", 2, 1],      # 인원 없음 → 성인+소인 3
    [4, "1인 조포", "가상손님D", "0304", "2026-10-20", "2026-10-22", 1, 0],  # 오늘 출발 → 내일 조식 없음
    [5, "1인 조포", "단체", "1030", "2026-10-22", "2026-10-23", 1, 0],      # 가상 객실 → 제외
]
RESERVE = [
    ["예약번호", "예약상태", "고객명", "객실번호", "도착일", "출발일", "객실수", "성인", "소인", "비고"],
    ["R1", "투숙", "가상손님C", "0303", "2026-10-22", "2026-10-23", 1, 2, 1, "조포"],  # 이미 체크인
    ["R2", "예약", "가상손님E", "0401", "2026-10-22", "2026-10-24", 1, 2, 0, "2인 조식 포함"],
    ["R3", "예약", "가상손님F", "", "2026-10-22", "2026-10-23", 1, 1, 0, "1인 조포"],  # 객실 미배정
    ["R4", "취소", "가상손님G", "", "2026-10-22", "2026-10-23", 1, 2, 0, "2인 조포"],
    ["R5", "예약", "가상손님H", "", "2026-10-23", "2026-10-24", 1, 2, 0, "2인 조포"],  # 내일 도착
    ["R6", "예약", "가상손님I", "", "2026-10-22", "2026-10-23", 1, 2, 0, "룸온리"],
    ["R7", "예약", "가상손님A", "0301", "2026-10-21", "2026-10-24", 1, 2, 0, "2인 조식 포함"],
]


def test_breakfast_count(cfg, tmp_path):
    """투숙 중 + 오늘 도착 예정(체크인 전)을 합쳐 내일 아침 조식 인원만 센다."""
    res = cv.breakfast_count(_xlsx(tmp_path / "ih.xlsx", INHOUSE), _xlsx(tmp_path / "rv.xlsx", RESERVE),
                             str(tmp_path / "out"), BDATE, cfg)
    assert res["bf_date"] == "2026-10-23"
    assert res["inhouse"] == 2 + 3          # 0301 2명 + 0303 3명 (0302 룸온리, 0304 오늘 출발, 1030 제외)
    assert res["arrivals"] == 2 + 1         # R2 2명 + R3 1명 (R1 체크인, R4 취소, R5 내일, R6 룸온리, R7 다른 날 도착)
    assert res["total"] == 8 and res["count"] == 4
    assert res["warnings"][0].startswith("10/23 아침 조식 8명")
    ws, v = _values(res["output"])
    assert ws["C2"].value == 8
    kinds = [r[0] for r in v[6:]]
    assert kinds.count("투숙 중") == 2 and kinds.count("도착 예정") == 2
    assert "미배정" in [r[1] for r in v[6:]]


def test_breakfast_count_no_status_column(cfg, tmp_path):
    """예약상태 열이 없어도 같은 객실·같은 이름이면 이미 체크인한 손님으로 보고 한 번만 센다."""
    rv = [[c for i, c in enumerate(r) if i != 1] for r in RESERVE]
    res = cv.breakfast_count(_xlsx(tmp_path / "ih.xlsx", INHOUSE), _xlsx(tmp_path / "rv.xlsx", rv),
                             str(tmp_path / "out"), BDATE, cfg)
    # 상태가 없으니 R4(취소)도 들어감 → 2+1+2
    assert res["inhouse"] == 5 and res["arrivals"] == 5
    assert any("예약상태" in w for w in res["warnings"])


def test_breakfast_count_ignores_zero_fee_column(cfg, tmp_path):
    """실제 객실 예약 리스트(2026-10-28)의 '조식요금[총액]' 열(모두 0)을 조식 여부로 읽어 0명이 되던 문제 → 비고로 판단."""
    rv = [["Acct#", "객실#", "도착일자", "출발일자", "Adult", "Child", "객실", "투숙객명", "상태", "비고", "조식요금[총액]"],
          ["A1", "", "2026-10-22", "2026-10-24", 1, 0, 1, "가상손님E", "확정", "1인조포\n\n현장결제예정", 0],
          ["A2", "", "2026-10-22", "2026-10-23", 1, 0, 1, "가상단체", "확정", "룸온니", 0]]
    res = cv.breakfast_count(_xlsx(tmp_path / "ih.xlsx", INHOUSE), _xlsx(tmp_path / "rv.xlsx", rv),
                             str(tmp_path / "out"), BDATE, cfg)
    assert res["arrivals"] == 1


def test_card_manual_rows_for_other_terminal(cfg, tmp_path):
    """카드: ERP와 연동 안 되는 한벽루 단말기 승인 건을 적는 노란 빈 줄 4개, 계에 포함. 승인 실패는 제외 (2026-10-28)."""
    res = cv.convert_file("card", _xlsx(tmp_path / "카드.xlsx", CARD), str(tmp_path / "out"), BDATE, cfg)
    ws, v = _values(res["output"])
    data = v[2:]
    assert [r[1] for r in data[:1]] == ["1111-XXXX-XXXX-XXXX"] and not any(r[5] == "실패" for r in data)
    manual = [i for i, r in enumerate(data) if r[8] == "한벽루 단말기 (직접 입력)"]
    assert len(manual) == 4 and all(not any(data[i][:8]) for i in manual)
    assert ws.cell(3 + manual[0], 8).fill.fgColor.rgb.endswith("FFF2B3")
    cash_row, total_row = 3 + manual[-1] + 1, 3 + manual[-1] + 2
    assert v[cash_row - 1][5] == "현금영수증"
    assert v[total_row - 1][7] == f"=SUM(H3:H{cash_row - 1})+H{cash_row}"  # 빈 줄까지 계에 포함


def test_card_room_fill_from_reservation_list(cfg, tmp_path):
    """카드 객실 자동 채우기: 객실 예약 리스트(CHECK IN)의 Acct# → 객실# (2026-10-28).
    한 Acct#에 객실이 여러 개면 승인금액 = 일일요금×일수인 객실, 그래도 모르면 후보를 모두 적고 노란 칸."""
    head = CARD[0]
    def card(acct, amt, no):
        return [BDATE, no, 0, acct, "카드승인", "1111-xxxx-xxxx-xxxx", 0, 0, amt, amt, "0000000000", "성공",
                "", "가상카드", "1", "", str(no), ""]
    cards = [head, card("A1", 198000, 1), card("A2", 210000, 2), card("A2", 250000, 3), card("A3", 440000, 4),
             card("A9", 130000, 5)]
    rv = [["NS", "Acct#", "객실#", "도착일자", "일수", "출발일자", "투숙객명", "일일요금", "상태"],
          [0, "A1", "0307", "2026-10-26", 2, "2026-10-28", "가상A", 99000, "CHECK IN"],
          [0, "A2", "0501", "2026-10-28", 2, "2026-10-30", "가상B", 105000, "CHECK IN"],
          [0, "A2", "0708", "2026-10-28", 2, "2026-10-30", "가상B", 125000, "CHECK IN"],
          [0, "A3", "0301", "2026-10-28", 2, "2026-10-30", "가상C", 110000, "CHECK IN"],
          [0, "A3", "0310", "2026-10-28", 2, "2026-10-30", "가상C", 110000, "CHECK IN"]]
    res = cv.convert_file("card", _xlsx(tmp_path / "c.xlsx", cards), str(tmp_path / "out"), BDATE, cfg,
                          room_source=_xlsx(tmp_path / "rv.xlsx", rv))
    ws, v = _values(res["output"])
    rows = [(r[0], r[7]) for r in v[2:7]]
    assert rows == [("301/310", 440000), ("307", 198000), ("501", 210000), ("708", 250000), (None, 130000)]
    assert ws["A3"].fill.fgColor.rgb.endswith("FFF2B3") and not ws["A4"].fill.fgColor.rgb.endswith("FFF2B3")
    assert any("A9" in w for w in res["warnings"]) and any("301/310" in w for w in res["warnings"])


def test_breakfast_group_total_remark_not_per_room(cfg, tmp_path):
    """단체 비고 '조식 포함 (35명)'은 단체 전체 인원 → 객실마다 35명으로 세지 않고 투숙객명 수로 센다 (2026-10-29: 632명 문제)."""
    rv = [["Acct#", "객실#", "도착일자", "출발일자", "Adult", "Child", "투숙객명", "상태", "비고"]]
    for i in range(3):
        rv.append(["G1", f"040{i + 1}", "2026-10-22", "2026-10-23", 1, 0, "가상갑/가상을", "확정", "조식 포함 (5명)\n현장 결제"])
    rv.append(["G1", "0404", "2026-10-22", "2026-10-23", 1, 0, "가상병", "확정", "조식 포함 (5명)\n현장 결제"])
    res = cv.breakfast_count(_xlsx(tmp_path / "ih.xlsx", INHOUSE), _xlsx(tmp_path / "rv.xlsx", rv),
                             str(tmp_path / "out"), BDATE, cfg)
    assert res["arrivals"] == 7  # 2+2+2+1 (비고의 5명은 단체 전체 — 여기선 이름 수 7과 다름 → 경고)
    assert any("단체 비고" in w and "다름" in w for w in res["warnings"])
    ok, n, why = cv._bf_people("조식 포함 (35명)", cfg["converters"]["breakfast"])
    assert ok and n is None and "단체 35명" in why
    assert cv._bf_people("2인 조식 포함", cfg["converters"]["breakfast"])[1] == 2


def test_mask_multiple_names():
    assert cv.mask_korean_name("홍길동/김철") == "홍0동/김0"
    assert cv.mask_korean_name("Zhu bao jun / Lin jia hao") == "Zhu bao jun / Lin jia hao"


def test_breakfast_count_makes_hotel_form_list(cfg, tmp_path, monkeypatch):
    """5번 조식 인원수: 투숙 중 + 도착 예정을 합친 호텔 양식 조식명단도 만든다 (2026-10-29). 미배정 객실은 맨 아래."""
    monkeypatch.setattr(cv, "app_dir", lambda: ROOT)
    res = cv.breakfast_count(_xlsx(tmp_path / "ih.xlsx", INHOUSE), _xlsx(tmp_path / "rv.xlsx", RESERVE),
                             str(tmp_path / "out"), BDATE, cfg)
    assert os.path.basename(res["list"]) == "조식명단(예약포함) 2026-10-23.xlsx" and res["sanity_ok"]
    ws, v = _values(res["list"])
    assert v[0][0] == "조 식 명 단(10/23)" and ws["B3"].font.name == "맑은 고딕"
    body = [r for r in v[2:] if r[0] and (str(r[0]).isdigit() or r[0] == "미배정")]
    assert [r[0] for r in body] == ["0301", "0303", "0401", "미배정"]
    total = next(r for r in v if r[0] == "합계")
    assert total[2].startswith("=SUM(")


def test_pax_sanity_warns_on_big_room(cfg):
    """안전장치: 한 객실 인원이 최대 인원(4)을 넘으면 맨 위에 '확인 필요' 경고 (632명 사고 재발 방지)."""
    w = ["요약"]
    assert not cv._pax_sanity([{"room": "0302", "pax": 35}, {"room": "0303", "pax": 2}], cfg["converters"]["breakfast"], w)
    assert w[0].startswith("⚠ 확인 필요") and "0302(35명)" in w[0]
    assert cv._pax_sanity([{"room": "0303", "pax": 4}], cfg["converters"]["breakfast"], [])


def test_find_daily_files(cfg, tmp_path):
    """오늘 마감 한 번에: ERP원본 폴더의 '약칭+MMDD' 파일 (수입1022, 카드1022 …), 날짜가 다른 파일은 무시 (2026-10-29)."""
    d = tmp_path / "ERP원본"
    d.mkdir()
    _xlsx(d / "수입1022.xlsx", INCOME)
    _xlsx(d / "수입1021.xlsx", INCOME)
    _xlsx(d / "카드1022.xlsx", CARD)
    _xlsx(d / "인하1022.xlsx", INHOUSE)
    _xlsx(d / "객실 1022.xlsx", RESERVE)
    (d / "~$수입1022.xlsx").write_bytes(b"")  # 엑셀 임시 파일은 무시
    jobs, notes, found = cv.daily_jobs(str(d), BDATE, cfg)
    assert {k: os.path.basename(v) for k, v in found.items()} == {
        "income": "수입1022.xlsx", "card": "카드1022.xlsx", "inhouse": "인하1022.xlsx", "rooms": "객실 1022.xlsx"}
    assert [(k, os.path.basename(p), os.path.basename(e) if e else None) for k, p, e in jobs] == [
        ("income", "수입1022.xlsx", None), ("card", "카드1022.xlsx", "객실 1022.xlsx"),
        ("bf_count", "인하1022.xlsx", "객실 1022.xlsx")]
    assert any("대외1022" in n for n in notes)


def test_breakfast_count_late_checkin_counted(cfg, tmp_path):
    """IN HOUSE LIST를 받은 뒤 체크인한 손님(객실조회에는 CHECK IN, IN HOUSE에는 없음)은 투숙 중으로 센다 (2026-10-29)."""
    rv = [["Acct#", "객실#", "도착일자", "출발일자", "Adult", "Child", "투숙객명", "상태", "비고"],
          ["A1", "0501", "2026-10-22", "2026-10-24", 1, 0, "가상갑", "CHECK IN", "1인 조포"],
          ["A2", "0301", "2026-10-21", "2026-10-24", 2, 0, "가상손님A", "CHECK IN", "2인 조식 포함"]]  # IN HOUSE에 있음
    res = cv.breakfast_count(_xlsx(tmp_path / "ih.xlsx", INHOUSE), _xlsx(tmp_path / "rv.xlsx", rv),
                             str(tmp_path / "out"), BDATE, cfg)
    assert res["inhouse"] == 5 + 1 and res["arrivals"] == 0
    assert any("뒤 체크인한 1객실" in w and "0501" in w for w in res["warnings"])
