# -*- coding: utf-8 -*-
"""거성호텔 나이트 마감 자동화 - 변환 로직 (GUI와 분리되어 단독 테스트 가능)"""
import csv
import datetime as dt
import json
import os
import re
import sys
from copy import copy
from html.parser import HTMLParser

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ConversionError(Exception):
    """사용자에게 그대로 보여 줄 수 있는 오류 메시지를 담는다."""


# ---------------------------------------------------------------- 설정 로드
def app_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def bundled_dir():
    return getattr(sys, "_MEIPASS", app_dir())


def load_config():
    """exe 옆의 config.json을 먼저 읽고, 없으면 exe 안에 포함된 기본값을 사용."""
    for base in (app_dir(), bundled_dir()):
        path = os.path.join(base, "config.json")
        if os.path.exists(path):
            with open(path, encoding="utf-8") as f:
                return json.load(f)
    raise ConversionError("config.json 파일을 찾을 수 없습니다.")


# ---------------------------------------------------------------- 파일 읽기
class _TableParser(HTMLParser):
    """ERP가 'xls'라는 이름으로 내보내는 HTML 표 파일을 읽기 위한 간단한 파서."""

    def __init__(self):
        super().__init__()
        self.tables, self._t, self._row, self._cell = [], None, None, None

    def handle_starttag(self, tag, attrs):
        if tag == "table" and self._t is None:
            self._t = []
            self.tables.append(self._t)
        elif tag == "tr" and self._t is not None:
            self._row = []
            self._t.append(self._row)
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None:
            self._row.append("".join(self._cell).strip())
            self._cell = None
        elif tag == "tr":
            self._row = None
        elif tag == "table":
            self._t = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _decode(raw):
    for enc in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def read_table(path):
    """엑셀(xlsx/xls), HTML 위장 xls, CSV를 모두 2차원 리스트로 읽는다."""
    if not os.path.exists(path):
        raise ConversionError(f"파일을 찾을 수 없습니다: {path}")
    with open(path, "rb") as f:
        raw = f.read()
    head = raw[:2048]

    if head[:2] == b"PK":  # xlsx / xlsm
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows

    if head[:4] == b"\xd0\xcf\x11\xe0":  # 구형 xls
        try:
            import xlrd
        except ImportError:
            raise ConversionError("구형 .xls 파일을 읽으려면 xlrd가 필요합니다. "
                                  "엑셀에서 .xlsx로 저장한 뒤 다시 시도해 주세요.")
        book = xlrd.open_workbook(path)
        sh = book.sheet_by_index(0)
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            rows.append(row)
        return rows

    text = _decode(raw)
    if head.lstrip()[:1] == b"<" or "<table" in text[:5000].lower():
        parser = _TableParser()
        parser.feed(text)
        if not parser.tables:
            raise ConversionError("파일에서 표를 찾지 못했습니다.")
        return max(parser.tables, key=len)

    dialect = csv.excel_tab if text.count("\t") > text.count(",") else csv.excel
    return [row for row in csv.reader(text.splitlines(), dialect)]


# ---------------------------------------------------------------- 값 정리
def norm(v):
    if v is None:
        return ""
    return re.sub(r"[\s_\-\(\)\[\]/\.\*:#]", "", str(v)).lower()


def clean_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def to_number(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    s = str(v).strip().replace(",", "").replace("원", "").replace("₩", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    try:
        n = float(s)
    except ValueError:
        return None
    n = -n if neg else n
    return int(n) if n.is_integer() else n


def to_date(v):
    if v is None or v == "":
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)) and 30000 < float(v) < 80000:  # 엑셀 일련번호
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).strftime("%Y-%m-%d")
    s = str(v).strip()
    m = re.search(r"(\d{4})[-./년\s]+(\d{1,2})[-./월\s]+(\d{1,2})", s)
    if m:
        return f"{int(m[1]):04d}-{int(m[2]):02d}-{int(m[3]):02d}"
    m = re.fullmatch(r"(\d{4})(\d{2})(\d{2})", s)
    if m:
        return f"{m[1]}-{m[2]}-{m[3]}"
    m = re.fullmatch(r"(\d{2})[-./](\d{1,2})[-./](\d{1,2})", s)
    if m:
        return f"20{m[1]}-{int(m[2]):02d}-{int(m[3]):02d}"
    return s


def mask_card(s):
    """마스킹되지 않은 카드번호만 앞4/뒤4자리만 남기고 가린다."""
    if "*" in s:
        return s
    digits = re.sub(r"\D", "", s)
    if len(digits) >= 13:
        return f"{digits[:4]}-****-****-{digits[-4:]}"
    return s


# ---------------------------------------------------------------- 헤더 탐지
def find_header(rows, fields, max_scan=40):
    """상단 max_scan행 중 별칭이 가장 많이 일치하는 행을 헤더로 판단한다.
    1단계: 정확히 일치, 2단계: 포함 관계(남은 열/항목만)."""
    alias_norm = {k: [norm(a) for a in v] for k, v in fields.items()}
    best = (0, None, {})
    for i, row in enumerate(rows[:max_scan]):
        mapping, used = {}, set()
        for j, cell in enumerate(row):
            n = norm(cell)
            if not n:
                continue
            for key, al in alias_norm.items():
                if key not in mapping and n in al:
                    mapping[key] = j
                    used.add(j)
                    break
        for j, cell in enumerate(row):
            n = norm(cell)
            if not n or j in used:
                continue
            for key, al in alias_norm.items():
                if key in mapping:
                    continue
                if any(len(a) >= 2 and a in n for a in al):
                    mapping[key] = j
                    used.add(j)
                    break
        if len(mapping) > best[0]:
            best = (len(mapping), i, mapping)
    return best


_TOTAL_WORDS = {"합계", "소계", "총계", "계", "total", "subtotal", "grandtotal"}


def _is_total_row(row):
    return any(norm(c) in _TOTAL_WORDS for c in row[:4])


def fmt_room(v, mode):
    s = clean_text(v)
    if re.fullmatch(r"\d{1,4}", s):
        return s.zfill(4) if mode == "pad4" else str(int(s))
    return s


def _room_key(rec):
    r = rec.get("room", "") or ""
    return (0, int(r)) if r.isdigit() else (1, r)


# ---------------------------------------------------------------- 원본 → 레코드
def _extract(conv, rows):
    """원본 표 → {field: value} 레코드 목록. (records, warnings, mapped_fields)"""
    fields = conv["source_fields"]
    score, hidx, mapping = find_header(rows, fields)
    warnings = []

    for tgt, alt in conv.get("fallback", {}).items():
        if tgt not in mapping and alt in mapping:
            mapping[tgt] = mapping[alt]
            warnings.append(f"'{fields[tgt][0]}' 열이 없어 '{fields[alt][0]}' 열로 대체했습니다.")

    missing_req = [k for k in conv["required"] if k not in mapping]
    if hidx is None or score == 0 or missing_req:
        names = ", ".join(fields[k][0] for k in missing_req) or "제목 행"
        found = [clean_text(c) for c in (rows[hidx] if hidx is not None else []) if clean_text(c)]
        raise ConversionError(
            f"원본에서 필수 항목({names})을 찾지 못했습니다.\n"
            f"선택한 파일이 맞는지 확인해 주세요.\n(인식된 제목: {', '.join(found[:12]) or '없음'})")

    for k in conv.get("warn_missing", []):
        if k not in mapping:
            warnings.append(f"원본에 '{fields[k][0]}' 열이 없어 빈칸으로 남깁니다.")

    types = conv.get("types", {})
    include = {f: {norm(v) for v in vals} for f, vals in conv.get("include", {}).items()}
    exclude = {f: {norm(v) for v in vals} for f, vals in conv.get("exclude", {}).items()}
    records, src_total = [], None
    for row in rows[hidx + 1:]:
        if not any(clean_text(c) for c in row):
            continue
        if _is_total_row(row):
            src_total = row
            continue
        rec = {}
        for key, j in mapping.items():
            val = row[j] if j < len(row) else None
            t = types.get(key)
            if t == "number":
                rec[key] = to_number(val)
            elif t == "date":
                rec[key] = to_date(val)
            else:
                rec[key] = clean_text(val)
        if any(f in rec and norm(rec[f]) not in vals for f, vals in include.items()):
            continue
        if any(f in rec and norm(rec[f]) in vals for f, vals in exclude.items()):
            continue
        if "room" in rec and conv.get("room_format"):
            rec["room"] = fmt_room(rec["room"], conv["room_format"])
        if conv.get("require_room") and not rec.get("room"):
            continue
        if not any(v not in ("", None) for v in rec.values()):
            continue
        records.append(rec)

    # 원본에 합계 행이 있으면 우리가 계산한 합계와 대조
    if conv.get("check_total") and src_total is not None and conv.get("sum"):
        fld = conv["sum"][0]
        j = mapping.get(fld)
        if j is not None and j < len(src_total):
            tv = to_number(src_total[j])
            mine = sum(r.get(fld) or 0 for r in records)
            if tv is not None and abs(tv - mine) > 0.5:
                warnings.append(f"원본 합계({tv:,.0f})와 변환 합계({mine:,.0f})가 다릅니다. 원본을 확인해 주세요.")
    return records, warnings, set(mapping)


def _is_breakfast(val, no_values):
    if val is None:
        return False
    n = norm(val)
    return n not in {norm(x) for x in no_values} and (to_number(val) is None or to_number(val) > 0)


# ---------------------------------------------------------------- 엑셀 출력
_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BLANK_FILL = PatternFill("solid", fgColor="FFF2B3")  # 근무자가 직접 확인/입력할 칸
_NO_FILL = PatternFill(fill_type=None)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_DEFAULT_FONT = {"name": "굴림", "title": 20, "header": 9, "body": 9}
_GREY = "666666"


def resolve_template(rel):
    """양식 파일은 exe 옆의 templates/를 먼저 찾고, 없으면 exe 안에 포함된 것을 찾는다."""
    if not rel:
        return None
    for base in (app_dir(), bundled_dir()):
        p = os.path.join(base, rel)
        if os.path.exists(p):
            return p
    return None


def _style_of(cell):
    return {"font": copy(cell.font), "border": copy(cell.border), "fill": copy(cell.fill),
            "alignment": copy(cell.alignment), "number_format": cell.number_format}


def _apply(cell, st, **over):
    st = {**st, **over}
    cell.font, cell.border, cell.fill = copy(st["font"]), copy(st["border"]), copy(st["fill"])
    cell.alignment = copy(st["alignment"])
    if st.get("number_format"):
        cell.number_format = st["number_format"]


def _has_style(cell):
    b = cell.border
    return any(getattr(b, s).style for s in ("left", "right", "top", "bottom")) or cell.fill.fill_type is not None


class _Sheet:
    """출력 시트와 서식 정보. 호텔 양식(템플릿)이 있으면 그 서식을, 없으면 config의 폰트를 쓴다."""

    def __init__(self):
        self.wb = self.ws = None
        self.hrow = 0
        self.colmap = {}        # out_cols 순번 → 시트 열 번호(1부터)
        self.data_st = {}       # 시트 열 → 데이터 칸 서식
        self.foot_st = {}       # 시트 열 → 합계 칸 서식
        self.data_height = None
        self.foot_height = None
        self.note_font = None
        self.from_template = False
        self.warnings = []


def _fill(color):
    return PatternFill("solid", fgColor=color) if color else _NO_FILL


def _new_sheet(sheet_name, title_text, out_cols, font_cfg, row_height):
    """템플릿이 없을 때: config의 font 설정(폰트·크기·행 높이)으로 새 양식을 만든다."""
    f = {**_DEFAULT_FONT, **(font_cfg or {})}
    s = _Sheet()
    s.wb = Workbook()
    ws = s.ws = s.wb.active
    ws.title = sheet_name[:31]
    nprint = sum(1 for c in out_cols if c.get("print", True))

    t = ws.cell(1, 1, title_text)
    t.font = Font(name=f["name"], size=f["title"], bold=f.get("title_bold", True))
    t.alignment = _CENTER
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nprint)
    ws.row_dimensions[1].height = f.get("title_height", 44)
    s.hrow = 2
    ws.row_dimensions[2].height = f.get("header_height", 28)
    head_fill = _fill(f.get("header_fill", "E7EBF0"))
    sum_fill = _fill(f.get("sum_fill", "F2F2F2"))
    for i, col in enumerate(out_cols):
        c = i + 1
        s.colmap[i] = c
        color = _GREY if not col.get("print", True) else "000000"
        cell = ws.cell(2, c, col["header"])
        cell.font = Font(name=f["name"], size=f["header"], bold=f.get("header_bold", True), color=color)
        cell.fill, cell.border, cell.alignment = copy(head_fill), _BORDER, _CENTER
        ws.column_dimensions[get_column_letter(c)].width = col.get("width", 12)
        s.data_st[c] = {"font": Font(name=f["name"], size=f["body"], bold=f.get("body_bold", False), color=color),
                        "border": _BORDER, "fill": _NO_FILL, "alignment": _CENTER, "number_format": None}
        s.foot_st[c] = {"font": Font(name=f["name"], size=f["body"], bold=f.get("sum_bold", True)),
                        "border": _BORDER, "fill": sum_fill, "alignment": _CENTER, "number_format": None}
    s.data_height = row_height or f.get("row_height")
    s.foot_height = s.data_height
    s.note_font = Font(name=f["name"], size=max(f["body"] - 1, 8), italic=True, color=_GREY)
    ws.freeze_panes = ws.cell(3, 1)
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "2:2"
    return s


def _template_sheet(path, sheet_name, title_text, out_cols):
    """호텔 양식을 열어 머리글 행을 찾고, 그 아래 데이터·합계 행의 서식을 기억한 뒤 비운다.
    폰트·열 너비·제목 병합·인쇄 설정은 양식 그대로 유지된다."""
    try:
        wb = load_workbook(path)
    except Exception as e:
        raise ConversionError(f"호텔 양식 파일을 열 수 없습니다: {os.path.basename(path)}\n({e})")
    ws = wb.worksheets[0]
    heads = {}  # 양식 머리글 → out_cols 순번 (header와 aliases 모두 인정)
    for i, col in enumerate(out_cols):
        for h in [col["header"]] + col.get("aliases", []):
            heads.setdefault(norm(h), i)

    best = (0, None, {})
    for r in range(1, min(ws.max_row, 30) + 1):
        m = {}
        for c in range(1, ws.max_column + 1):
            i = heads.get(norm(ws.cell(r, c).value))
            if i is not None and i not in m:
                m[i] = c
        if len(m) > best[0]:
            best = (len(m), r, m)
    if best[0] < 2:
        raise ConversionError(f"호텔 양식({os.path.basename(path)})에서 머리글 행(객실 등)을 찾지 못했습니다.\n"
                              "양식의 열 제목이 config.json의 output 'header'와 같은지 확인해 주세요.")
    s = _Sheet()
    s.wb, s.ws, s.from_template = wb, ws, True
    s.hrow, s.colmap = best[1], dict(best[2])
    hrow = s.hrow

    # 양식에 없는 열(예: 인쇄 안 되는 Acct# 참고 열)은 오른쪽 끝에 덧붙인다
    nxt = max(max(s.colmap.values()), ws.max_column) + 1
    for i, col in enumerate(out_cols):
        if i in s.colmap:
            continue
        s.colmap[i] = nxt
        src = ws.cell(hrow, max(best[2].values()))
        cell = ws.cell(hrow, nxt, col["header"])
        _apply(cell, _style_of(src))
        cell.font = copy(src.font) if col.get("print", True) else Font(
            name=src.font.name, size=src.font.size, bold=src.font.bold, color=_GREY)
        ws.column_dimensions[get_column_letter(nxt)].width = col.get("width", 12)
        if col.get("print", True):
            s.warnings.append(f"호텔 양식에 '{col['header']}' 열이 없어 오른쪽 끝에 추가했습니다.")
        nxt += 1

    # 제목: 머리글 위쪽의 첫 글자 칸을 이번 날짜 제목으로 교체(서식은 그대로)
    done = False
    for r in range(1, hrow):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip():
                ws.cell(r, c).value = title_text
                done = True
                break
        if done:
            break
    if not done:
        s.warnings.append("호텔 양식에서 제목 칸을 찾지 못해 제목을 넣지 않았습니다.")

    # 데이터 행 서식: 머리글 바로 아래 행. 서식이 없으면 머리글 서식에서 굵게만 뺀다
    drow = hrow + 1
    use_data = drow <= ws.max_row and any(_has_style(ws.cell(drow, c)) for c in best[2].values())
    foot_row = None
    for r in range(hrow + 1, ws.max_row + 1):
        if any(norm(ws.cell(r, c).value) in ("계", "합계", "총계") for c in range(1, ws.max_column + 1)):
            foot_row = r
            break
    if foot_row == drow:
        use_data = False
    for i, c in s.colmap.items():
        if use_data:
            st = _style_of(ws.cell(drow, c))
        else:
            st = _style_of(ws.cell(hrow, c))
            st["font"] = copy(st["font"])
            st["font"].b = False
            st["fill"] = _NO_FILL
        if not out_cols[i].get("print", True):
            fnt = copy(st["font"])
            fnt.color = _GREY
            st["font"] = fnt
        s.data_st[c] = st
        if foot_row:
            s.foot_st[c] = _style_of(ws.cell(foot_row, c))
        else:
            fst = dict(st)
            fnt = copy(st["font"])
            fnt.b = True
            fst["font"] = fnt
            s.foot_st[c] = fst
    rd = ws.row_dimensions
    s.data_height = rd[drow].height if use_data and drow in rd else None
    s.foot_height = rd[foot_row].height if foot_row and foot_row in rd else s.data_height
    base_font = s.data_st[s.colmap[0]]["font"]
    s.note_font = Font(name=base_font.name, size=base_font.size, italic=True, color=_GREY)

    # 머리글 아래(예시 행·빈 행·기존 합계 행)를 비운다
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row > hrow:
            ws.unmerge_cells(str(mr))
    if ws.max_row > hrow:
        ws.delete_rows(hrow + 1, ws.max_row - hrow)
    for k in [k for k in rd if k > hrow]:
        del rd[k]
    if not ws.print_title_rows:
        ws.print_title_rows = f"{hrow}:{hrow}"
    ws.title = sheet_name[:31]
    return s


def write_workbook(path, sheet_name, title_text, out_cols, records, sum_fields,
                   number_format="0", notes=None, cash=None, row_height=None,
                   template=None, font_cfg=None):
    """template(호텔 양식 경로)가 있으면 양식을 열어 값만 채우고, 없으면 font_cfg로 새로 만든다.
    반환: 경고 목록"""
    if template:
        s = _template_sheet(template, sheet_name, title_text, out_cols)
        if row_height and s.data_height is None:
            s.data_height = row_height
    else:
        s = _new_sheet(sheet_name, title_text, out_cols, font_cfg, row_height)
    ws, hrow, colmap = s.ws, s.hrow, s.colmap
    key = lambda c: c.get("field") or c.get("formula_of") or c["header"]
    idx = {key(c): i for i, c in enumerate(out_cols)}
    printable = sorted(colmap[i] for i, c in enumerate(out_cols) if c.get("print", True))
    letter = lambda i: get_column_letter(colmap[i])

    r = hrow
    for rec in records:
        r += 1
        if s.data_height:
            ws.row_dimensions[r].height = s.data_height
        for i, col in enumerate(out_cols):
            cell = ws.cell(r, colmap[i])
            st = s.data_st[colmap[i]]
            _apply(cell, st)
            fmt = st["number_format"] if st["number_format"] not in (None, "General") else number_format
            if col.get("formula_of"):
                cell.value = f"={letter(idx[col['formula_of']])}{r}"
                cell.number_format = fmt
                continue
            f = col.get("field")
            val = rec.get(f) if f else None
            if val not in ("", None):
                cell.value = val
            if f is None and col.get("fill", True):
                cell.fill = copy(_BLANK_FILL)
            if isinstance(val, (int, float)):
                cell.number_format = fmt

    first, last = hrow + 1, r

    def _foot_row(row):
        if s.foot_height:
            ws.row_dimensions[row].height = s.foot_height
        for c in printable:
            _apply(ws.cell(row, c), s.foot_st[c])

    def _foot(row, i, value, fill=None):
        cell = ws.cell(row, colmap[i])
        cell.value = value
        if fill is not None:
            cell.fill = copy(fill)
        if not isinstance(value, str) or value.startswith("="):
            fmt = s.foot_st[colmap[i]]["number_format"]
            cell.number_format = fmt if fmt not in (None, "General") else number_format

    if records and cash is not None:  # 카드 양식: 현금영수증 줄 + 계
        r += 1
        _foot_row(r)
        _foot(r, idx["kind"], "현금영수증")
        amt = idx["amount"]
        _foot(r, amt, cash["amount"], None if cash["amount"] is not None else _BLANK_FILL)
        cash_row = r
        r += 1
        _foot_row(r)
        L = letter(amt)
        _foot(r, idx["approval_no"], "계")
        _foot(r, amt, f"=SUM({L}{first}:{L}{last})+{L}{cash_row}")
    elif records and sum_fields:
        r += 1
        _foot_row(r)
        sum_idx = [idx[f] for f in sum_fields if f in idx]
        if sum_idx:
            label_i = min(sum_idx, key=lambda i: colmap[i])
            label_col = max(colmap[label_i] - 1, printable[0])
            lab = ws.cell(r, label_col)
            lab.value = "계"
            for i in sum_idx:
                L = letter(i)
                _foot(r, i, f"=SUM({L}{first}:{L}{last})")

    if notes:
        r += 2
        for n in notes:
            ws.cell(r, printable[0], n).font = copy(s.note_font)
            r += 1

    ws.print_area = f"A1:{get_column_letter(printable[-1])}{r}"

    try:
        s.wb.save(path)
    except PermissionError:
        raise ConversionError(f"결과 파일이 엑셀에서 열려 있습니다. 닫은 뒤 다시 시도해 주세요.\n{path}")
    return s.warnings


# ---------------------------------------------------------------- 공개 API
def unique_path(path):
    base, ext = os.path.splitext(path)
    n, p = 2, path
    while os.path.exists(p):
        p = f"{base}({n}){ext}"
        n += 1
    return p


def _parse_date(s):
    try:
        return dt.datetime.strptime(s, "%Y-%m-%d").date()
    except ValueError:
        raise ConversionError("영업일자를 YYYY-MM-DD 형식으로 입력해 주세요. 예) 2026-10-23")


def _bf_people(remark, conv):
    """비고 문구로 조식 대상 여부와 조식 인원을 판단. (대상여부, 인원 또는 None)"""
    t = remark or ""
    if re.search(conv["bf_exclude_regex"], t, re.I):
        return False, None
    if re.search(conv["bf_include_regex"], t, re.I):
        m = re.search(r"(\d+)\s*인", t)
        return True, (int(m.group(1)) if m else None)
    return False, None


def convert_file(kind, src_path, out_dir, business_date, cfg=None, unique=False):
    """kind: income | card | credit | breakfast.
    반환: {'output': 경로, 'count': 건수, 'warnings': [...], 'title': 이름}"""
    cfg = cfg or load_config()
    conv = cfg["converters"][kind]
    bdate = _parse_date(business_date)
    rows = read_table(src_path)
    records, warnings, mapped = _extract(conv, rows)
    notes = []
    cash = None

    if kind == "income":
        warnings.insert(0, "매출 열은 ERP의 NET(공급가, 부가세 제외) 금액으로 채웠습니다. 호텔 양식 기준과 같은지 확인해 주세요.")
        warnings.append("'선/후', '숙박수'는 노란 칸에 직접 입력하세요. '누계'는 우선 매출과 같은 값(수식)으로 채워져 있습니다.")

    if kind == "card":
        if conv.get("cash_footer"):
            kws = conv.get("cash_keywords", ["현금"])
            cash_rows = [r for r in records if any(k in (r.get("kind") or "") for k in kws)]
            records = [r for r in records if r not in cash_rows]
            total_cash = sum(r.get("amount") or 0 for r in cash_rows)
            cash = {"amount": total_cash if cash_rows else None}
            if cash_rows:
                warnings.append(f"현금영수증 {len(cash_rows)}건({total_cash:,.0f}원)은 표에서 빼고 '현금영수증' 줄에 합산했습니다.")
            else:
                warnings.append("원본에 현금영수증 건이 없습니다. '현금영수증' 줄(노란 칸)은 필요하면 직접 입력하세요.")
        if cfg.get("mask_card_number", True):
            for rec in records:
                if rec.get("card_no"):
                    rec["card_no"] = mask_card(rec["card_no"])
        for rec in records:
            if rec.get("card_no"):
                rec["card_no"] = re.sub(r"x", "X", rec["card_no"], flags=re.I) if "x" in rec["card_no"].lower() else rec["card_no"]
        if "room" not in mapped:
            warnings.insert(0, "ERP 승인내역에는 객실번호가 없고 Acct#만 있습니다. '객실' 열은 비워 두었으니 "
                               "오른쪽 'Acct#' 열을 참고해 직접 입력해 주세요. (Acct#는 인쇄되지 않습니다)")

    if kind == "breakfast":
        has_col = "breakfast" in mapped and any(r.get("breakfast") not in ("", None) for r in records)
        before = len(records)
        kept = []
        for rec in records:
            if has_col:
                ok, n = _is_breakfast(rec.get("breakfast"), conv["no_breakfast_values"]), None
            else:
                ok, n = _bf_people(rec.get("remark"), conv)
            if not ok:
                continue
            if n is not None:
                rec["pax"] = n
            elif rec.get("pax") is None:
                a, c = rec.get("adults"), rec.get("children")
                if a is not None or c is not None:
                    rec["pax"] = (a or 0) + (c or 0)
            kept.append(rec)
        records = kept
        notes.append(f"※ 전체 {before}객실 중 조식 대상 {len(records)}객실 (비고 문구 기준 자동 추출)")
        no_pax = sum(1 for r in records if r.get("pax") is None)
        if no_pax:
            warnings.append(f"조식 대상 중 {no_pax}객실은 인원을 알 수 없어 빈칸입니다. 직접 입력해 주세요.")
        if "remark" not in mapped and not has_col:
            warnings.append("원본에 '비고'(조식 표시) 열이 없어 조식 대상을 가려내지 못했습니다. 전체를 확인해 주세요.")
        warnings.append("비고에 '조식'/'조포'가 있는 객실만 뽑았습니다('룸온리'는 제외). 비고가 다른 표현이면 알려 주세요.")

    if conv.get("sort_by_room"):
        records.sort(key=_room_key)

    if not records:
        raise ConversionError("변환할 데이터 행이 없습니다. 파일 내용이나 조건(조식 대상 등)을 확인해 주세요.")

    shown = bdate + dt.timedelta(days=conv.get("date_offset_days", 0))
    md = shown.strftime("%m/%d")
    title_text = conv["sheet_title"].format(md=md)
    fname = conv["file_pattern"].format(date=shown.strftime("%Y-%m-%d"), m=shown.month, d=shown.day)
    fname = re.sub(r'[\\/:*?"<>|]', "-", fname)

    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"{fname}.xlsx")
    if unique:
        out_path = unique_path(out_path)
    tpl_rel = conv.get("template")
    tpl = resolve_template(tpl_rel)
    if tpl_rel and not tpl:
        warnings.append(f"호텔 양식 파일({tpl_rel})이 없어 config.json의 기본 서식으로 만들었습니다.")
    warnings += write_workbook(out_path, shown.strftime("%m%d"), title_text, conv["output"], records,
                               conv.get("sum", []), conv.get("number_format", "0"), notes, cash,
                               conv.get("row_height"), template=tpl, font_cfg=conv.get("font"))
    return {"output": out_path, "count": len(records), "warnings": warnings, "title": conv["title"]}


def detect_type(path, cfg=None):
    """파일명 키워드 → 제목 행 고유 항목 순으로 마감 서류 종류를 추정. 실패 시 None."""
    cfg = cfg or load_config()
    name = os.path.basename(path).lower()
    for key, conv in cfg["converters"].items():
        if any(kw.lower() in name for kw in conv["keywords"]):
            return key
    try:
        rows = read_table(path)
    except Exception:
        return None
    best, best_key = (0, 0), None
    for key, conv in cfg["converters"].items():
        score, _, mapping = find_header(rows, conv["source_fields"])
        sig = sum(1 for s in conv.get("signature", []) if s in mapping)
        if sig > 0 and (sig, score) > best:
            best, best_key = (sig, score), key
    return best_key


def default_out_dir(src_path):
    return os.path.join(os.path.dirname(os.path.abspath(src_path)), "마감결과")


def default_business_date(cfg=None, now=None):
    """나이트 근무는 자정을 넘기므로, 기준 시각(기본 12시) 이전이면 전날을 영업일자로 본다."""
    cfg = cfg or {}
    now = now or dt.datetime.now()
    cut = cfg.get("business_day_rollover_hour", 12)
    d = now.date() - dt.timedelta(days=1) if now.hour < cut else now.date()
    return d.strftime("%Y-%m-%d")


def today_str():
    return dt.date.today().strftime("%Y-%m-%d")
